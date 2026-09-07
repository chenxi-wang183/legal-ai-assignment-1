
import os
 
# --- NLTK DATA LOCATION (must run BEFORE importing llama_index) ---
# LlamaIndex ships an nltk cache inside site-packages whose files are hardlinked.
# Some sandboxed hosts refuse to open multiply-linked files, which breaks indexing
# with a "Security Violation [pathsec.open]" error. Pointing NLTK_DATA at a normal
# writable directory makes nltk find its data there first and never touch that cache.
NLTK_DIR = "/tmp/nltk_data"
os.environ["NLTK_DATA"] = NLTK_DIR
os.makedirs(NLTK_DIR, exist_ok=True)
try:
    import nltk
    nltk.data.path.insert(0, NLTK_DIR)
    for _pkg in ("stopwords", "punkt", "punkt_tab"):
        try:
            nltk.download(_pkg, download_dir=NLTK_DIR, quiet=True)
        except Exception:
            pass
except Exception:
    pass
 
import streamlit as st
import tempfile
from llama_index.core import VectorStoreIndex, Document, Settings
from llama_index.llms.openai_like import OpenAILike
from llama_index.embeddings.openai_like import OpenAILikeEmbedding
from llama_index.core.node_parser import SentenceSplitter, TokenTextSplitter
from pathlib import Path
 
# --- PAGE CONFIGURATION (Must be the first command) ---
st.set_page_config(page_title="Legal AI Assistant", layout="wide")
 
 
# --- MODEL PROVIDER CONFIGURATION ---
# All providers below expose an OpenAI-compatible API, so the same client code
# works for each one - only the endpoint and the model names change.
PROVIDERS = {
    "Zhipu GLM (free)": {
        "api_base": "https://open.bigmodel.cn/api/paas/v4",
        "llm_model": "glm-4-flash",
        "embed_model": "embedding-3",
        "context_window": 128000,
        "key_hint": "open.bigmodel.cn",
    },
    "Google Gemini": {
        "api_base": "https://generativelanguage.googleapis.com/v1beta/openai/",
        "llm_model": "gemini-2.0-flash",
        "embed_model": "text-embedding-004",
        "context_window": 1000000,
        "key_hint": "aistudio.google.com",
    },
    "OpenAI": {
        "api_base": "https://api.openai.com/v1",
        "llm_model": "gpt-4o-mini",
        "embed_model": "text-embedding-3-small",
        "context_window": 128000,
        "key_hint": "platform.openai.com",
    },
}
 
 
# --- CUSTOM STYLES & ICONS ---
scales_svg = """
<svg width="80" height="80" viewBox="0 0 24 24" fill="none" xmlns="http://www.w3.org/2000/svg">
<path d="M21 6L18.5 5M18.5 5L15 6M18.5 5V12M18.5 12L21 13M18.5 12L15 13M3 6L5.5 5M5.5 5L9 6M5.5 5V12M5.5 12L3 13M5.5 12L9 13M12 3V21M3 21H21" stroke="#1e3a8a" stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round"/>
</svg>
"""
sources_svg = """
<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="#2563eb" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
<path d="M2 3h6a4 4 0 0 1 4 4v14a3 3 0 0 0-3-3H2z"></path>
<path d="M22 3h-6a4 4 0 0 0-4 4v14a3 3 0 0 1 3-3h7z"></path>
</svg>
"""
custom_css = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Cormorant+Garamond:wght@300;400;600&display=swap');
    .stApp { background-color: #ffffff; }
    .main-container { display: flex; flex-direction: column; align-items: center; padding: 1rem 1rem 5rem 1rem; }
    .title { font-family: 'Cormorant Garamond', serif; font-weight: 600; font-size: 5rem; color: #1e3a8a; text-align: center; margin-top: 2rem; margin-bottom: 0rem; }
    .subtitle { font-family: 'Cormorant Garamond', serif; font-weight: 400; color: #9ca3af; text-align: center; margin-top: 3rem; margin-bottom: 3rem; font-size: 1.25rem; line-height: 1.6; }
    [data-testid="stHorizontalBlock"] { align-items: center; }
    .stTextInput input { background-color: #ffffff !important; border: 1px solid #d1d5db; height: 3.5rem; border-radius: 0.5rem; padding-left: 1rem; }
    [data-testid="stFileUploader"] { width: 3.5rem; height: 3.5rem; }
    [data-testid="stFileUploader"] section { border: none; padding: 0; }
    [data-testid="stFileUploader"] section > input + div { display: none; }
    [data-testid="stFileUploader"] [data-testid="stFileUploaderFile"] { display: none; }
    [data-testid="stFileUploader"] section button { display: flex; justify-content: center; align-items: center; background-color: transparent !important; border: none !important; color: #9ca3af !important; padding: 0 !important; width: 3.5rem; height: 3.5rem; font-size: 0; }
    [data-testid="stFileUploader"] section button:hover { background-color: #eef2ff !important; color: #2563eb !important; }
    [data-testid="stFileUploader"] section button::after { content: '+'; font-size: 2.5rem; font-weight: 300; line-height: 1; }
    .stButton>button { background-color: transparent !important; color: #9ca3af !important; border: none !important; width: 3.5rem !important; height: 3.5rem !important; font-size: 1.5rem !important; padding: 0 !important; margin: 0 !important; }
    .stButton>button:hover { background-color: #eef2ff !important; color: #2563eb !important; }
    .response-container { font-size: 1.1rem; line-height: 1.7; }
    .stSidebar { background-color: #ffffff; }
    blockquote { background-color: #f3f4f6; border-left: 5px solid #d1d5db; padding: 1rem; border-radius: 0.25rem; color: #4b5563; }
</style>
"""
st.markdown(custom_css, unsafe_allow_html=True)
 
 
# --- SYSTEM PROMPT ---
system_prompt = (
    "You are a legal AI assistant. Your ONLY function is to answer questions based on provided text. "
    "It is absolutely CRUCIAL that you follow the specified four-step format for EVERY response, without exception. "
    "This format is MANDATORY for all questions, including simple questions or complex summarization tasks.\n\n"
    "--- RESPONSE FORMAT ---\n\n"
    "[ANALYSIS]\n"
    "Break down the user's question into its core components.\n\n"
    "[RELEVANT_CLAUSES]\n"
    "Cite the exact clause number(s) and quote the original text verbatim.\n\n"
    "[DIRECT_ANSWER]\n"
    "Provide a one-sentence summary answer. If the question is a yes/no question, start with "
    "'Yes' or 'No' and then state the core reason. If it is an open question (for example asking "
    "what a document covers, or asking you to summarise), give the summary directly instead of "
    "forcing a Yes or No.\n\n"
    "[REASONING]\n"
    "Explain how the cited clauses logically lead to your direct answer.\n\n"
    "--- END OF FORMAT ---\n\n"
    "If the provided text does not contain the answer, you MUST ignore the format and ONLY state: "
    "'Based on the provided document, I could not find a specific answer to this question.'"
)
 
# --- SESSION STATE INITIALIZATION ---
if "rag_engine" not in st.session_state: st.session_state.rag_engine = None
if "history" not in st.session_state: st.session_state.history = []
if "analysis_result" not in st.session_state: st.session_state.analysis_result = None
if "api_key" not in st.session_state: st.session_state.api_key = ""
if "provider" not in st.session_state: st.session_state.provider = ""
if "last_uploaded_filename" not in st.session_state: st.session_state.last_uploaded_filename = None
 
 
def extract_documents(file_path: str, file_name: str):
    """Extract plain text ourselves.
 
    SimpleDirectoryReader relies on pypdf, which returns nothing for some PDFs and
    attaches a large block of XMP metadata (pdf:Keywords, pdfx:*) to every document.
    When extraction fails, that metadata becomes the only content in the index and
    every answer ends up quoting it. Extracting page by page avoids both problems.
    """
    docs = []
    lower = file_name.lower()
 
    if lower.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        text=text,
                        metadata={"page_label": str(page_number), "file_name": file_name},
                    ))
    elif lower.endswith(".docx"):
        from docx import Document as DocxDocument
        text = "\n".join(p.text for p in DocxDocument(file_path).paragraphs)
        if text.strip():
            docs.append(Document(text=text, metadata={"file_name": file_name}))
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            docs.append(Document(text=text, metadata={"file_name": file_name}))
 
    # Metadata is for citing sources only - it should never be embedded or shown
    # to the model, otherwise it competes with the actual clauses during retrieval.
    for doc in docs:
        doc.excluded_embed_metadata_keys = ["page_label", "file_name"]
        doc.excluded_llm_metadata_keys = ["page_label", "file_name"]
 
    return docs
 
 
def build_index(documents):
    """Build the vector index, falling back to a splitter that does not use nltk."""
    try:
        splitter = SentenceSplitter(chunk_size=512, chunk_overlap=50)
        return VectorStoreIndex.from_documents(documents, transformations=[splitter])
    except Exception:
        # SentenceSplitter relies on nltk data, which some hosts block. TokenTextSplitter
        # splits on plain separators instead, so it always works - slightly rougher chunks.
        splitter = TokenTextSplitter(chunk_size=512, chunk_overlap=50)
        return VectorStoreIndex.from_documents(documents, transformations=[splitter])
 
 
def configure_models(provider_name: str, api_key: str) -> None:
    """Point LlamaIndex at the selected OpenAI-compatible endpoint."""
    cfg = PROVIDERS[provider_name]
 
    # Some LlamaIndex internals still read this variable, so keep it in sync.
    os.environ["OPENAI_API_KEY"] = api_key
 
    # OpenAILike is used instead of OpenAI because the OpenAI class validates the
    # model name against a hard-coded list and rejects anything it does not know.
    Settings.llm = OpenAILike(
        model=cfg["llm_model"],
        api_base=cfg["api_base"],
        api_key=api_key,
        is_chat_model=True,          # required, otherwise the completion endpoint is called
        context_window=cfg["context_window"],
        temperature=0.1,
        timeout=120,
        system_prompt=system_prompt,
    )
    Settings.embed_model = OpenAILikeEmbedding(
        model_name=cfg["embed_model"],
        api_base=cfg["api_base"],
        api_key=api_key,
        embed_batch_size=10,
    )
 
 
# --- SIDEBAR ---
with st.sidebar:
    st.header("Model Provider")
    provider_name = st.selectbox(
        "Provider", list(PROVIDERS.keys()), key="provider_select", label_visibility="collapsed"
    )
    st.caption(f"Get a key at {PROVIDERS[provider_name]['key_hint']}")
 
    st.header("API Credentials")
    api_key_input = st.text_input(
        "Enter your API Key", type="password", key="api_key_input_sidebar",
        label_visibility="collapsed", placeholder="Enter your API Key..."
    )
 
    key_changed = api_key_input and api_key_input != st.session_state.api_key
    provider_changed = provider_name != st.session_state.provider
 
    if api_key_input and (key_changed or provider_changed):
        try:
            configure_models(provider_name, api_key_input)
            st.session_state.api_key = api_key_input
            st.session_state.provider = provider_name
            st.success(f"✅ Connected to {provider_name}")
        except Exception as e:
            st.error(f"Failed to configure models. Please check your key. Error: {e}")
            st.session_state.api_key = ""  # Invalidate key on error
            st.session_state.provider = ""
 
        # Embedding dimensions differ between providers, so any existing index
        # must be discarded and the document re-indexed.
        st.session_state.rag_engine = None
        st.session_state.last_uploaded_filename = None
        st.session_state.analysis_result = None
 
    st.header("Search History")
    if st.button("Clear History"):
        st.session_state.history = []
        st.rerun()
 
    if not st.session_state.history:
        st.write("No searches yet.")
    else:
        for i, item in enumerate(reversed(st.session_state.history)):
            with st.expander(f"**{i+1}. {item['question'][:50]}...**"):
                st.markdown(item.get('formatted_answer', '...'), unsafe_allow_html=True)
 
 
# --- MAIN PAGE UI ---
st.markdown('<div class="main-container">', unsafe_allow_html=True)
st.markdown(f'<div style="text-align: center;">{scales_svg}</div>', unsafe_allow_html=True)
st.markdown('<p class="title">Legal AI Assistant</p>', unsafe_allow_html=True)
st.markdown('<p class="subtitle">Upload a legal document.<br>Ask a question.<br>Get a structured analysis.</p>', unsafe_allow_html=True)
 
if not st.session_state.api_key:
    st.warning("Please choose a provider and enter your API Key in the sidebar to begin.")
else:
    # --- SEARCH BAR LAYOUT ---
    search_bar_cols = st.columns([1, 8, 1])
    with search_bar_cols[0]:
        uploaded_file = st.file_uploader("Upload", type=["pdf", "txt", "docx"], key="file_uploader", label_visibility="collapsed")
    with search_bar_cols[1]:
        user_question = st.text_input("Enter your question...", key="question_input", label_visibility="collapsed")
    with search_bar_cols[2]:
        analyze_button_clicked = st.button("➤", key="analyze_button", help="Analyze the document", use_container_width=True)
 
    # --- DOCUMENT INDEXING ---
    if uploaded_file:
        if "last_uploaded_filename" not in st.session_state or st.session_state.last_uploaded_filename != uploaded_file.name:
            with st.spinner("Indexing the document..."):
                try:
                    with tempfile.TemporaryDirectory() as temp_dir:
                        temp_file_path = os.path.join(temp_dir, uploaded_file.name)
                        with open(temp_file_path, "wb") as f: f.write(uploaded_file.getbuffer())
 
                        documents = extract_documents(temp_file_path, uploaded_file.name)
 
                        if not documents:
                            st.error(
                                "No text could be extracted from this file. It may be a scanned "
                                "PDF (images only), which would need OCR."
                            )
                            st.stop()
 
                        total_chars = sum(len(d.text) for d in documents)
                        index = build_index(documents)
                        st.session_state.rag_engine = index.as_query_engine(similarity_top_k=8)
                        st.session_state.last_uploaded_filename = uploaded_file.name
                    st.success(f"✅ Indexed {len(documents)} sections, {total_chars:,} characters extracted.")
                    st.session_state.analysis_result = None
                except Exception as e:
                    st.error(f"Indexing failed: {e}")
 
    # --- QUERY LOGIC ---
    if analyze_button_clicked:
        if st.session_state.rag_engine and user_question:
            with st.spinner("Thinking..."):
                try:
                    response_obj = st.session_state.rag_engine.query(user_question)
                    response_text = str(response_obj)
 
                    parts = response_text.split('[')
                    analysis, clauses, answer, reasoning = "", "", "", ""
                    for part in parts:
                        if part.startswith("ANALYSIS]"): analysis = part.replace("ANALYSIS]", "").strip()
                        elif part.startswith("RELEVANT_CLAUSES]"): clauses = part.replace("RELEVANT_CLAUSES]", "").strip()
                        elif part.startswith("DIRECT_ANSWER]"): answer = part.replace("DIRECT_ANSWER]", "").strip()
                        elif part.startswith("REASONING]"): reasoning = part.replace("REASONING]", "").strip()
 
                    formatted_answer = ""
                    if analysis and clauses and answer and reasoning:
                        formatted_answer += f"<blockquote><b>Analysis & Reasoning:</b><br>{analysis}<br><br>{reasoning}</blockquote><hr>"
                        formatted_answer += f"<b>Direct Answer:</b><br>{answer}<br><hr>"
                        formatted_answer += f"<b>Relevant Legal Clause(s):</b><br>{clauses}"
                    else:
                        formatted_answer = response_text
 
                    st.session_state.analysis_result = (response_obj, formatted_answer)
 
                    if not any(d['question'] == user_question for d in st.session_state.history):
                        st.session_state.history.append({'question': user_question, 'formatted_answer': formatted_answer})
 
                except Exception as e:
                    st.error(f"An error occurred: {e}")
        elif not st.session_state.rag_engine:
             st.warning("⚠️ Please upload a document first.")
        else:
            st.warning("⚠️ Please enter a question.")
 
    # --- DISPLAY RESPONSE ---
    if st.session_state.analysis_result:
        response_obj, formatted_answer = st.session_state.analysis_result
        st.write("---")
        st.markdown('<div class="response-container">', unsafe_allow_html=True)
        st.markdown(formatted_answer, unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)
 
        expander_title = f'<div style="display: flex; align-items: center; gap: 10px;">{sources_svg}<span>Show Cited Sources</span></div>'
        st.markdown(expander_title, unsafe_allow_html=True)
        with st.expander(" ", expanded=False):
            for i, node in enumerate(response_obj.source_nodes):
                with st.container(border=True):
                    page_label = node.metadata.get('page_label')
                    if page_label:
                        st.markdown(f"**Source from Page: {page_label}** (Similarity: {node.score:.4f})")
                    else:
                        st.markdown(f"**Source {i+1}** (Similarity: {node.score:.4f})")
                    st.write(node.get_text())
 
st.markdown('</div>', unsafe_allow_html=True)
 









